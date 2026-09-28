"""Workday adapter: course list + weekly meeting-pattern schedule, read from
a student's own "View My Courses" export (.xlsx). There's no session here -
Workday is downloaded by the student themselves in a browser and handed to
`lauds sync workday --file <path>`; see lauds.session for adapters that keep
a logged-in session instead.

Clean port of main's hub/workday.py - same behaviour, split so the parsing
is pure over raw bytes (parity tests replay the committed fixtures through
it) and `fetch()` only does the "read the file" part:

- A real export has one row per meeting *component* (Lecture, Lab, ...),
  each repeating the same "Course Listing" text. `parse_courses` dedupes to
  one Course per code (first occurrence wins); `parse_schedule` keeps one
  Meeting per component row instead.
- Two real-export hazards get the same defenses as main (#49 / e3db118):
  a declared `<dimension>` that understates the real row range (streamed
  with `reset_dimensions()` instead of trusted), and a stray populated cell
  at Excel's max address that would otherwise make a naive full-grid read
  hang for over a minute (bounded by the same sane row/col ceilings as
  main and `web/lib/workday.js`).
- A line in "Meeting Patterns" that doesn't parse (an async/online
  component with no fixed time, a bare "-", a TBD day or time) is skipped,
  never guessed at.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

import openpyxl

from lauds.models import Bundle, Course, Meeting, normalise_course_code

NAME = "workday"
DESCRIPTION = "Workday: course list and weekly class schedule, from a downloaded export"

_SHEET_NAME = "View My Courses"
_HEADER_ALIASES = {
    "course": {"course listing", "course"},
    "meeting_patterns": {"meeting patterns"},
    "instructional_format": {"instructional format"},
}
# ponytail: fixed ceilings, same values as main (web/lib/workday.js's twin) -
# bound a pathological file, not a real export. Upgrade: a size-based cutoff
# if a real export ever needs more than this.
_MAX_SANE_ROW = 10_000
_MAX_SANE_COL = 500

_DAY_CODES = {"mon": "MO", "tue": "TU", "wed": "WE", "thu": "TH", "fri": "FR", "sat": "SA", "sun": "SU"}

# Workday's own "Instructional Format" values -> a Meeting.kind. An unlisted
# format (a component type we haven't seen) falls back to "class", the same
# "safest generic bucket" idea as Item.kind -> category_for's default.
_KIND_FOR_FORMAT = {
    "lecture": "lecture",
    "laboratory": "lab",
    "seminar": "seminar",
    "tutorial": "tutorial",
    "discussion": "tutorial",
    "exam": "exam",
    "final exam": "exam",
}

_TIME_RE = re.compile(r"(\d{1,2}:\d{2})\s*([ap])\.?m\.?", re.IGNORECASE)


def _find_header_row(rows):
    """(row_index, {logical name: column index}) for the row holding a
    Course Listing/Course header; the other two logical names are read off
    that same row when present, and simply absent from the dict on an older
    export that lacks them - not an error, so parse_schedule degrades to "no
    meetings" instead of crashing."""
    for row_index, row in enumerate(rows):
        found = {}
        for col_index, cell in enumerate(row):
            value = str(cell).strip().lower() if cell is not None else ""
            for name, aliases in _HEADER_ALIASES.items():
                if value in aliases:
                    found[name] = col_index
        if "course" in found:
            return row_index, found
    return None, {}


def _cell(row, columns, name):
    col = columns.get(name)
    if col is None or col >= len(row) or row[col] is None:
        return None
    return str(row[col]).strip()


def _course_from_listing(listing, term):
    code_part, _, title = listing.partition(" - ")
    faculty, number, section = normalise_course_code(code_part)
    if not faculty or not number:
        return None
    return Course(code=f"{faculty} {number}", section=section, term=term,
                  title=title.strip() or listing, source=NAME)


def _read_rows(raw: bytes):
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook[_SHEET_NAME] if _SHEET_NAME in workbook.sheetnames else workbook.active
        sheet.reset_dimensions()
        return [row[:_MAX_SANE_COL + 1] for row in sheet.iter_rows(max_row=_MAX_SANE_ROW + 1, values_only=True)]
    finally:
        workbook.close()


def parse_courses(raw: bytes, term: str) -> list[Course]:
    """A "View My Courses" export's raw bytes -> one Course per unique code."""
    rows = _read_rows(raw)
    header_row, columns = _find_header_row(rows)
    if header_row is None:
        return []
    courses, seen = [], set()
    for row in rows[header_row + 1:]:
        listing = _cell(row, columns, "course")
        if not listing:
            continue
        course = _course_from_listing(listing, term)
        if course is not None and course.code not in seen:
            seen.add(course.code)
            courses.append(course)
    return courses


def _parse_time(text):
    m = _TIME_RE.match(text.strip())
    if not m:
        raise ValueError(f"unrecognized time: {text!r}")
    return datetime.strptime(f"{m.group(1)} {m.group(2).upper()}M", "%I:%M %p").time()


def _parse_pattern(pattern, course_code, kind, source):
    """One "Meeting Patterns" cell holds one line per meeting component, e.g.
    "2026-09-08 - 2026-12-05 | Mon Wed Fri | 10:00 a.m. - 11:00 a.m. | Building
    | Floor: 1 | Room: 100". A line can also be a bare "-" or something
    unparseable (an async component, a TBD exam slot) - skip that line
    rather than guess at a time."""
    meetings = []
    for line in pattern.splitlines():
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 3:
            continue
        date_range, days_text, times_text, *location_parts = parts
        try:
            start_str, end_str = (d.strip() for d in date_range.split(" - "))
            term_start = datetime.strptime(start_str, "%Y-%m-%d").date()
            term_end = datetime.strptime(end_str, "%Y-%m-%d").date()
        except ValueError:
            continue
        days = [_DAY_CODES[d[:3].lower()] for d in days_text.split() if d[:3].lower() in _DAY_CODES]
        if not days:
            continue
        try:
            start_text, end_text = (t.strip() for t in times_text.split(" - "))
            start_time = _parse_time(start_text)
            end_time = _parse_time(end_text)
        except ValueError:
            continue
        meetings.append(Meeting(
            course=course_code, kind=kind, days=days, start_time=start_time, end_time=end_time,
            location=", ".join(p for p in location_parts if p),
            term_start=term_start, term_end=term_end, source=source,
        ))
    return meetings


def parse_schedule(raw: bytes, term: str, source: str = NAME) -> list[Meeting]:
    """Same export's raw bytes -> one Meeting per meeting-component row.
    Returns [] (never raises) on an older export missing the Meeting
    Patterns column, or a row whose pattern doesn't parse."""
    rows = _read_rows(raw)
    header_row, columns = _find_header_row(rows)
    if header_row is None or "meeting_patterns" not in columns:
        return []
    meetings = []
    for row in rows[header_row + 1:]:
        listing = _cell(row, columns, "course")
        pattern = _cell(row, columns, "meeting_patterns")
        if not listing or not pattern:
            continue
        course = _course_from_listing(listing, term)
        if course is None:
            continue
        format_text = (_cell(row, columns, "instructional_format") or "").lower()
        kind = _KIND_FOR_FORMAT.get(format_text, "class")
        meetings += _parse_pattern(pattern, course.code, kind, source)
    return meetings


def fetch(path, term: str, source: str = NAME) -> Bundle:
    """`path` is a downloaded "View My Courses" export sitting on disk - the
    only "network" here is the student exporting it themselves in a
    browser, out of scope for this adapter (no login, no session)."""
    with open(path, "rb") as f:
        raw = f.read()
    return Bundle(courses=parse_courses(raw, term), meetings=parse_schedule(raw, term, source))
