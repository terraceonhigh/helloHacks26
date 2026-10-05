import re
from datetime import datetime

import openpyxl

from hub.logic import normalise_course_code
from hub.models import Course, Meeting

_SHEET_NAME = "View My Courses"
_COURSE_HEADER_ALIASES = {"course listing", "course"}
# Zero-based ceilings, same values as web/lib/workday.js's MAX_SANE_ROW/COL.
# A real export is a few hundred rows; these only bound pathological files.
_MAX_SANE_ROW = 10_000
_MAX_SANE_COL = 500

# Other columns on the same sheet/row hub/workday.py's course parser already
# reads - both ignored until now. Same file, same header row, no new
# fixture: see fixtures/workday_view_my_courses.xlsx.
_HEADER_ALIASES = {
    "course": _COURSE_HEADER_ALIASES,
    "meeting_patterns": {"meeting patterns"},
    "instructional_format": {"instructional format"},
}


def _find_header_row(rows):
    """(row_index, {logical name: column index}) for the row containing a
    "Course Listing"/"Course" header - the other two are read from that same
    row if present, and simply absent from the dict (not an error) on an
    older export that doesn't have them, so parse_workday_schedule degrades
    to "no meetings found" instead of crashing."""
    for row_index, row in enumerate(rows):
        found = {}
        for col_index, cell in enumerate(row):
            value = str(cell).strip().lower() if cell is not None else ""
            for name, aliases in _HEADER_ALIASES.items():
                if value in aliases:
                    found[name] = col_index
        if "course" in found:
            return row_index, found
    return None, {}


def _cell(row, columns, name):
    col = columns.get(name)
    if col is None or col >= len(row) or row[col] is None:
        return None
    return str(row[col]).strip()


def _course_from_listing(listing, term):
    code_part, _, title = listing.partition(" - ")
    faculty, number, section = normalise_course_code(code_part)
    if not faculty or not number:
        return None

    return Course(
        code=f"{faculty} {number}",
        section=section,
        term=term,
        title=title.strip() or listing,
    )


def _read_rows(path):
    # Two real-export hazards, handled together (#49):
    # - the declared <dimension> understates the populated range, so a
    #   read_only sheet that trusts it silently truncates. reset_dimensions()
    #   drops the declared range and streams the cells actually present.
    # - a stray cell at Excel's max address (XFD1048576) made read_only=False
    #   iter_rows() materialise the full 1M x 16K grid (>60s, >2GB).
    #   Streaming + the same sane ceilings web/lib/workday.js uses keep it
    #   bounded.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[_SHEET_NAME] if _SHEET_NAME in workbook.sheetnames else workbook.active
        sheet.reset_dimensions()
        return [row[:_MAX_SANE_COL + 1] for row in sheet.iter_rows(max_row=_MAX_SANE_ROW + 1, values_only=True)]
    finally:
        workbook.close()


def parse_workday_courses(path, term):
    rows = _read_rows(path)
    header_row_index, columns = _find_header_row(rows)
    if header_row_index is None:
        return []

    # Real exports have one row per meeting component (Lecture, Lab,
    # Discussion...), each repeating the same course listing text - dedupe
    # by code, keeping the first occurrence.
    courses = []
    seen_codes = set()
    for row in rows[header_row_index + 1:]:
        listing = _cell(row, columns, "course")
        if not listing:
            continue
        course = _course_from_listing(listing, term)
        if course is not None and course.code not in seen_codes:
            seen_codes.add(course.code)
            courses.append(course)
    return courses


_DAY_CODES = {"mon": "MO", "tue": "TU", "wed": "WE", "thu": "TH", "fri": "FR", "sat": "SA", "sun": "SU"}

# Workday's own "Instructional Format" values -> a Meeting.kind. Unlisted
# formats (a component type we haven't seen) fall back to "class" - the
# safest generic bucket, same idea as Item's kind->category default.
_KIND_FOR_FORMAT = {
    "lecture": "lecture",
    "laboratory": "lab",
    "seminar": "seminar",
    "tutorial": "tutorial",
    "discussion": "tutorial",
    "exam": "exam",
    "final exam": "exam",
}

_TIME_RE = re.compile(r"(\d{1,2}:\d{2})\s*([ap])\.?m\.?", re.IGNORECASE)


def _parse_time(text):
    m = _TIME_RE.match(text.strip())
    if not m:
        raise ValueError(f"unrecognized time: {text!r}")
    return datetime.strptime(f"{m.group(1)} {m.group(2).upper()}M", "%I:%M %p").time()


def _parse_meeting_pattern(pattern, course_code, kind, source):
    """One "Meeting Patterns" cell holds one line per meeting component,
    each shaped like:
    "2026-09-08 - 2026-12-05 | Mon Wed Fri | 10:00 a.m. - 11:00 a.m. | Building | Floor: 1 | Room: 100"
    A line can also be a bare "-" or something unparseable (an async/online
    component with no fixed time, a TBD exam slot) - skip that line rather
    than guess at a time."""
    meetings = []
    for line in pattern.splitlines():
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 3:
            continue
        date_range, days_text, times_text, *location_parts = parts
        try:
            start_str, end_str = (d.strip() for d in date_range.split(" - "))
            term_start = datetime.strptime(start_str, "%Y-%m-%d").date()
            term_end = datetime.strptime(end_str, "%Y-%m-%d").date()
        except ValueError:
            continue
        days = [_DAY_CODES[d[:3].lower()] for d in days_text.split() if d[:3].lower() in _DAY_CODES]
        if not days:
            continue
        try:
            start_text, end_text = (t.strip() for t in times_text.split(" - "))
            start_time = _parse_time(start_text)
            end_time = _parse_time(end_text)
        except ValueError:
            continue
        meetings.append(Meeting(
            course=course_code,
            kind=kind,
            days=days,
            start_time=start_time,
            end_time=end_time,
            location=", ".join(p for p in location_parts if p),
            term_start=term_start,
            term_end=term_end,
            source=source,
        ))
    return meetings


def parse_workday_schedule(path, term, source="workday"):
    """Recurring class meetings (lecture/lab/... times, days, location) from
    the same export parse_workday_courses reads - the "Meeting Patterns" and
    "Instructional Format" columns, both ignored until now. Returns []
    (never raises) on an older export missing those columns, or a row with
    no parseable meeting pattern - same "degrade, don't crash" contract as
    parse_workday_courses."""
    rows = _read_rows(path)
    header_row_index, columns = _find_header_row(rows)
    if header_row_index is None or "meeting_patterns" not in columns:
        return []

    meetings = []
    for row in rows[header_row_index + 1:]:
        listing = _cell(row, columns, "course")
        pattern = _cell(row, columns, "meeting_patterns")
        if not listing or not pattern:
            continue
        course = _course_from_listing(listing, term)
        if course is None:
            continue
        format_text = (_cell(row, columns, "instructional_format") or "").lower()
        kind = _KIND_FOR_FORMAT.get(format_text, "class")
        meetings += _parse_meeting_pattern(pattern, course.code, kind, source)
    return meetings
